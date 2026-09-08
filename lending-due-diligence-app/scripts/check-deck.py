"""
Checks the generated deck for layout problems that only show up when rendered:
shapes past the slide edge, text that will not fit its box, and tables whose
wrapped rows grow past the bottom of the slide.

Estimates are deliberately conservative (they over-predict height slightly), so
a clean run means real PowerPoint rendering has headroom.

Usage: python scripts/check-deck.py <deck.pptx>
"""

import sys
from pptx import Presentation
from pptx.util import Emu

EMU_IN = 914400
SAFE_BOTTOM = 7.30  # leave room for the page number

deck = sys.argv[1] if len(sys.argv) > 1 else "docs/Pega-to-PowerPlatform-Migration-PoC.pptx"
p = Presentation(deck)
SW, SH = p.slide_width / EMU_IN, p.slide_height / EMU_IN

# Segoe UI averages ~0.50 * point size per char in width at these sizes.
CHAR_W = 0.50
LINE_H = 1.32   # line height as a multiple of font size


def est_lines(text, width_in, size_pt):
    if not text:
        return 1
    chars_per_line = max(8, int(width_in * 72.0 / (size_pt * CHAR_W)))
    total = 0
    for seg in text.split("\n"):
        total += max(1, -(-len(seg) // chars_per_line))
    return total


def frame_height(tf, width_in):
    """Estimated rendered height of a text frame, in inches."""
    h = 0.0
    for para in tf.paragraphs:
        text = "".join(r.text for r in para.runs)
        size = 14.0
        for r in para.runs:
            if r.font.size:
                size = r.font.size.pt
                break
        lines = est_lines(text, width_in, size)
        h += lines * size * LINE_H / 72.0
        h += (para.space_after.pt if para.space_after else 0) / 72.0
        h += (para.space_before.pt if para.space_before else 0) / 72.0
    return h


problems = []
warnings = []

for i, slide in enumerate(p.slides, 1):
    for sh in slide.shapes:
        if sh.top is None or sh.height is None:
            continue
        left = sh.left / EMU_IN
        top = sh.top / EMU_IN
        w = sh.width / EMU_IN
        h = sh.height / EMU_IN

        if top + h > SH + 0.02:
            problems.append(f"slide {i}: shape extends to {top + h:.2f}in (slide is {SH:.2f}in)")
        if left + w > SW + 0.02:
            problems.append(f"slide {i}: shape extends right to {left + w:.2f}in (slide is {SW:.2f}in)")

        if sh.has_text_frame and sh.text_frame.text.strip():
            need = frame_height(sh.text_frame, w - 0.12)
            if top + need > SAFE_BOTTOM:
                problems.append(
                    f"slide {i}: text needs {need:.2f}in from y={top:.2f} "
                    f"-> ends {top + need:.2f}in, past {SAFE_BOTTOM}in "
                    f"[{sh.text_frame.text.strip().splitlines()[0][:48]!r}]")
            elif need > h + 0.05:
                warnings.append(
                    f"slide {i}: text ({need:.2f}in) taller than its box ({h:.2f}in) "
                    f"but still on-slide")

        if sh.has_table:
            tbl = sh.table
            cols = [c.width / EMU_IN for c in tbl.columns]
            total = 0.0
            for r_i, row in enumerate(tbl.rows):
                declared = row.height / EMU_IN if row.height else 0.30
                tallest = declared
                for c_i, cell in enumerate(row.cells):
                    txt = cell.text
                    size = 11.5
                    for para in cell.text_frame.paragraphs:
                        for run in para.runs:
                            if run.font.size:
                                size = run.font.size.pt
                                break
                        break
                    lines = est_lines(txt, cols[c_i] - 0.20, size)
                    need = lines * size * LINE_H / 72.0 + 0.08
                    tallest = max(tallest, need)
                total += tallest
            bottom = top + total
            if bottom > SAFE_BOTTOM:
                problems.append(
                    f"slide {i}: table grows to {bottom:.2f}in "
                    f"(from y={top:.2f}, {len(tbl.rows)} rows), past {SAFE_BOTTOM}in")

    # Shape-to-shape collisions: a table or box silently drawn over a
    # neighbouring one is invisible to bounds checks but obvious on screen.
    boxes = []
    for sh in slide.shapes:
        if sh.top is None or sh.height is None or sh.width is None:
            continue
        has_content = (sh.has_text_frame and sh.text_frame.text.strip()) or sh.has_table
        if not has_content:
            continue  # background bands and rules are meant to sit behind things
        w_in = sh.width / EMU_IN
        if w_in > 12.5:
            continue  # full-bleed banners
        boxes.append((sh.left / EMU_IN, sh.top / EMU_IN, w_in,
                      sh.height / EMU_IN, sh.has_table))

    for a in range(len(boxes)):
        for b in range(a + 1, len(boxes)):
            ax, ay, aw, ah, at = boxes[a]
            bx, by, bw, bh, bt = boxes[b]
            if not (at or bt):
                continue  # only flag overlaps involving a table
            ox = min(ax + aw, bx + bw) - max(ax, bx)
            oy = min(ay + ah, by + bh) - max(ay, by)
            if ox > 0.06 and oy > 0.06:
                problems.append(
                    f"slide {i}: shapes overlap by {ox:.2f} x {oy:.2f}in "
                    f"(x {max(ax, bx):.2f}, y {max(ay, by):.2f})")

print(f"{deck}\n{len(p.slides)} slides at {SW:.2f} x {SH:.2f} in\n")
if problems:
    print(f"PROBLEMS ({len(problems)}):")
    for x in problems:
        print("  " + x)
else:
    print("No layout problems detected.")
if warnings:
    print(f"\nNotes ({len(warnings)}):")
    for x in warnings:
        print("  " + x)

sys.exit(1 if problems else 0)
